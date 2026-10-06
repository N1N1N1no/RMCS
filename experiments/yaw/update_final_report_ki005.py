#!/usr/bin/env python3
"""Replace the historical PID figure with the completed Ki=0.05 A/B diagnostic."""

from pathlib import Path

from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / '云台系统辨识与控制优化总报告_2026-10-02.docx'
FIGURE = ROOT / 'current_ki005_pid_diagnostic.png'


def paragraph(document, prefix):
    found = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(found) != 1:
        raise ValueError(f'Expected one paragraph beginning {prefix!r}, got {len(found)}')
    return found[0]


def replace(document, prefix, new_text):
    p = paragraph(document, prefix)
    p.text = new_text
    return p


def main():
    doc = Document(REPORT)
    replace(doc, '本报告面向',
            '本报告面向 2027 赛季算法组控制方向“云台系统辨识与控制优化”考核，汇总 C 车 yaw 轴的辨识、控制器调优和直接对比。辨识、选参、前馈分离和扫频数据保留其采集时的配置；“当前最终控制器”的效果只依据速度环 Ki=0.05 部署后完成的同轨迹直接 A/B 与独立 Ki 配对试验。原始 CSV、复算脚本和统计文件列于文末。')
    replace(doc, '范围说明：',
            '范围说明：第 2–7 节及第 9 节记录历史阶段试验，相关曲线和数字不能视为 Ki=0.05 下重测；第 8 节为当前 Ki=0.05 最终控制器与原控制器的 ±0.10 rad 平滑轨迹直接 A/B，第 10 节为只改变 Ki 的低速配对 A/B。模型拟合优度不能直接代表控制器跟踪精度。模型前馈目前只在手动 yaw sweep 中生效；Ki=0.05 已在 C 车普通控制配置中启用，日常遥控与自瞄仍需单独验证。')
    replace(doc, '4 控制器改变后的模型复核', '4 历史控制器调整后的本体模型复核')
    replace(doc, '最终 PID 与前馈启用后',
            '在选定角度/速度环 Kp=12/11 并加入前馈之后、Ki 提高到 0.05 之前，做了三轮直接力矩试验。辨识服务在试验期间覆盖发送给 yaw 电机的控制器输出，因此比较的仍是本体输入力矩与输出角度，而不是闭环目标跟随。三轮都在预设的 0.12 rad 转角保护处提前结束；为了公平，新旧数据统一只取完整的第 1–2 个四相周期。此项验证不能替代当前 Ki=0.05 工作点的再次辨识。')
    replace(doc, '5 PID 基准与参数选择', '5 历史 PID 选参与当前参数复核')
    replace(doc, '第一阶段使用',
            '下表是历史阶跃选参记录：目标为 ±0.10 rad，速度环 Ki=0.02，未启用额外前馈或最终力矩限幅。它说明为何选择角度/速度 Kp=12/11，但不是 Ki=0.05 的阶跃复测。图 4 改用已完成的当前 Ki=0.05 平滑轨迹 A/B 原始数据，展示当前参数下的目标跟踪及坐标系差异。')
    replace(doc, '12/11 相对原',
            '在历史阶跃筛选中，12/11 相对原 10/13 的两次重复使 90% 到位时间缩短约 17%，IAE 降低约 10%，反馈力矩 P99 增加约 4%。阶跃会使控制器请求力矩峰值超过 20 N·m，因此后续改用 0.5 s 平滑轨迹，并在反馈与前馈相加后增加 4.0 N·m 命令限幅。这些历史阶跃指标不能当作当前 Ki=0.05 的实测指标。')
    replace(doc, '图 4  原 PID',
            '图 4  当前 Ki=0.05 与原控制器的同轨迹 A/B 第一对。上图的 yaw 由目标减去记录的控制角误差重构，并非独立传感器测量；下图由目标、相对底盘的编码器角度和控制角误差推得坐标系差量。')

    old_figure = paragraph(doc, '图 4  当前 Ki=0.05')._p.getprevious()
    if old_figure is None or not old_figure.xpath('.//a:blip'):
        raise ValueError('Figure 4 image not found immediately before its caption')
    from docx.text.paragraph import Paragraph
    figure_paragraph = Paragraph(old_figure, doc)
    figure_paragraph.clear()
    figure_paragraph.add_run().add_picture(str(FIGURE), width=Inches(6.32))

    caption = paragraph(doc, '图 4  当前 Ki=0.05')
    interpretation = doc.add_paragraph(
        '当前配置三轮的最后保持窗口（14.65–14.90 s）中，若把 Odom 坐标系目标直接与相对底盘的编码器角度相减，会看到 0.972–1.043° 的表观差值；同一窗口记录的控制角误差中位数只有 0.043–0.070°。两者的差主要来自参考坐标系不一致及底盘/零点变化，不能直接作为 PID 稳态误差。重构量只用于解释该差值，当前跟踪性能以控制角误差和第 8 节三对 A/B 为准。'
    )
    caption._p.addnext(interpretation._p)

    replace(doc, '保留 PID 12/11，分别测试',
            '历史前馈分离试验保留 PID 12/11、速度环 Ki=0.02，分别测试无前馈、仅加入目标角速度、再加入模型力矩。目标角速度参考为 ωd；模型力矩项为 τff = Bωd + Jαd。每组各运行一轮约 15 s，在相同轨迹上统计 7 次完整换向。本节数字不是 Ki=0.05 的复测结果。')
    replace(doc, '为检查单次前馈结果的重复性',
            '在 Ki=0.02 的历史配置下，为检查单次前馈结果的重复性，对三组候选 B、J 各做“当前→候选”三对交错测试，共 18 轮、54,005 行遥测。9 次当时参数重复的转动 RMS 为 0.0842 ± 0.0107°；IAE 为 (0.545 ± 0.052)×10⁻³ rad·s/次。这里的“当前”仅指当时的对照参数，不能当作现用 Ki=0.05 的重复数据。')
    replace(doc, '为检验原二阶本体模型能否外推',
            '在 Ki=0.05 调整前，为检验原二阶本体模型能否外推，做了 13 轮 8 s 线性扫频，共 99,358 行 CSV。命令由模型前馈加弱位置/速度校正产生，验证输入取实测电机力矩；原 J=0.14261 kg·m²、B=1.37855 N·m·s/rad 和 3 ms 延迟固定不变。每轮只用最初 1 s 估计常值扰动，其余数据计算加速度留出集 R²。低频组逐级加幅，并保留 0.18 rad 角度、1.2 rad/s 速度和 3 N·m 力矩保护。这是本体模型的历史外推证据，不是现用 Ki=0.05 控制器的跟踪测试。')
    replace(doc, '第一，原二阶本体模型',
            '第一，原二阶本体模型在高频小幅数据中较准，在 0.3–0.65 Hz 明显失配；该扫频采于 Ki=0.05 调整前，不能据此声称当前参数下模型全频段有效。第二，不同批次起始 yaw 约为 1.2–2.6 rad，低速 Ki A/B 期间起点也逐渐变化；交错和反序只能减轻顺序影响。第三，Ki=0.05 的低速收益已重复，较快轨迹只有两对，其 IAE 略增，日常遥控和自瞄仍需单独验证。第四，新一轮“原→当前最终”A/B 的 96% 以上收益是 Kp、Ki 和前馈组合后的效果，不能单独归因于 Ki；积分改动的增量由第 10 节单独评估。第五，模型前馈目前只由手动 yaw sweep 生成目标速度和加速度。第六，图 4 上编码器与目标的表观差值不等于实际控制角误差；其坐标系差量由现有记录推算，未用独立底盘姿态传感器复核。')

    index = doc.tables[-1]
    for row in index.rows:
        if row.cells[0].text == 'PID 调优':
            row.cells[0].text = '历史 PID 选参与当前复核'
            row.cells[2].text = 'pid_comparison.png；compare_pid.py；current_ki005_pid_diagnostic.png；current_ki005_pid_diagnostic.json；plot_current_ki005_pid_diagnostic.py'
        if row.cells[0].text == '当前最终配置直接 A/B':
            row.cells[2].text += '；current_ki005_pid_diagnostic.png；plot_current_ki005_pid_diagnostic.py'

    doc.save(REPORT)
    print(REPORT)


if __name__ == '__main__':
    main()
